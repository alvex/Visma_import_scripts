#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sammanstall_fakturor.py

Sammanställer alla betalningsposter som förberetts för import under en
kalendermånad till en kontrollista i Excel:

    <output-mapp>\fakturor_lista\fakturor_<mappnamn>.xlsx

där <mappnamn> är månadsmappens namn (arbetsmappen som dashboard.py arbetar
mot, t.ex. "sept 01"). Registrerings-CSV:erna ligger i månadsmappens
undermapp "output", så mappnamnet hämtas från output-mappens förälder.

Källfiler (läses, ändras aldrig):
    betalningar_lista_to_reg_*.csv              (Bankgiro)
    swish_betalningar_for_registrering_*.csv    (Swish)

Filer med "_visma" i namnet samt "~$"-temporärfiler exkluderas.

Syftet är en kontrollista: alla poster bevaras, återkommande fakturanummer
markeras som "Dubblett" (på samtliga berörda rader), men ingenting tas bort
eller slås ihop automatiskt. "Förberedd" betyder bara att posten finns i
underlaget för import – inte att importen är genomförd.

Omkörningsskydd: varje källfil identifieras med SHA-256 i arket "Källfiler".
En oförändrad fil som redan lästs in hoppas över (inga extra rader skapas,
även om filen innehåller två identiska rader – de bevaras som två rader från
första inläsningen). En ändrad källfil rapporteras tydligt och läses INTE in
igen; befintliga data skrivs aldrig över i tysthet.

Körs från terminalen:
    python sammanstall_fakturor.py --input "C:\\ekonomi\\Fak_2026\\sept 01\\output"
    python sammanstall_fakturor.py --input "C:\\ekonomi\\Fak_2026\\sept 01" --month 2026-10
    python sammanstall_fakturor.py --input "...\\output" --dry-run

Returkoder: 0 = OK, 1 = fel (inget skrevs), 2 = OK men ändrade källfiler
upptäcktes (kräver manuell kontroll).
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import os
import re
import sys
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("FEL: Paketet 'openpyxl' saknas. Installera med:  pip install openpyxl")
    sys.exit(1)


# --- Konfiguration -----------------------------------------------------------

# Prefix för registrerings-CSV:erna som ska läsas in.
SOURCE_PREFIXES = ("betalningar_lista_to_reg", "swish_betalningar_for_registrering")

# Undermapp och filnamnsprefix för sammanställningen.
TARGET_SUBDIR = "fakturor_lista"
TARGET_PREFIX = "fakturor"

# Kolumner som MÅSTE finnas i käll-CSV:erna (Betalningsreferens används inte).
SOURCE_REQUIRED_COLUMNS = ["Datum", "Avsändare", "Fakturanummer", "Belopp"]

# Arknamn och kolumner i sammanställningen.
SHEET_FAKTUROR = "Fakturor"
SHEET_FEL = "Felrader"
SHEET_KALLFILER = "Källfiler"
SHEET_META = "Metadata"

FAKTUROR_COLUMNS = ["Datum", "Avsändare", "Fakturanummer", "Belopp", "Status", "Source"]
FEL_COLUMNS = ["Källfil", "Rad", "Problem", "Datum", "Avsändare", "Fakturanummer", "Belopp"]
KALLFIL_COLUMNS = ["Källfil", "SHA256", "Inläst", "Period", "Rader", "Giltiga", "Felrader", "Utanför period"]

STATUS_PREPARED = "Förberedd"
STATUS_DUPLICATE = "Dubblett"

MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")


# --- Normalisering och tolkning ----------------------------------------------

def normalize_invoice(value: Any) -> str:
    """Normalisera ett fakturanummer till ren text.

    Tar bort alla blanksteg (även hårda mellanslag) men bevarar inledande
    nollor – därför hanteras heltal via str() och aldrig via någon
    numerisk omtolkning av strängar.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    text = str(value).replace("\u00a0", " ").replace("\u202f", " ")
    return re.sub(r"\s+", "", text)


def normalize_sender(value: Any) -> str:
    """Trimma avsändarnamn och normalisera inre blanksteg."""
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").replace("\u202f", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_date(value: Any) -> str | None:
    """Tolka ett datumvärde till ISO-text (ÅÅÅÅ-MM-DD). None om ogiltigt.

    Dagens datum används ALDRIG som ersättning för saknat datum.
    """
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value).strip()
    if not DATE_RE.match(text):
        return None
    try:
        return dt.date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def parse_amount(value: Any) -> Decimal | None:
    """Tolka ett belopp till Decimal med två decimaler. None om ogiltigt.

    Hanterar svenska format ("1 234,50", hårda mellanslag som tusental,
    decimalkomma) och standardformat ("1234.50", "1,234.50"). Negativa
    belopp bevaras med sitt tecken.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != value:  # NaN
            return None
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    text = str(value).strip()
    # Hårda mellanslag och vanliga mellanslag är tusentalsavgränsare.
    text = text.replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    # Typografiska minustecken -> vanligt minus.
    text = text.replace("\u2212", "-").replace("\u2013", "-")
    if not text:
        return None
    if "," in text and "." in text:
        # Sista förekomsten avgör vilket som är decimaltecken.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def format_sum_sv(total: Decimal) -> str:
    """Formatera en summa med svenskt format, t.ex. 41 899,00."""
    q = total.quantize(Decimal("0.01"))
    return f"{q:,.2f}".replace(",", " ").replace(".", ",")


def sha256_of_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# --- Inläsning av käll-CSV ----------------------------------------------------

def find_source_files(csv_dir: Path) -> list[Path]:
    """Hitta registrerings-CSV:erna i mappen. Endast avsedda källfiler."""
    files: list[Path] = []
    for prefix in SOURCE_PREFIXES:
        for p in sorted(csv_dir.glob(f"{prefix}_*.csv")):
            if not p.is_file():
                continue
            if p.name.startswith("~$"):
                continue
            if "_visma" in p.name.lower():
                continue
            files.append(p)
    return files


def read_source_csv(path: Path) -> tuple[list[dict], list[dict], int]:
    """Läs en registrerings-CSV.

    Returnerar (giltiga_rader, felrader, antal_datarader). Varje giltig rad:
    {"datum", "avsandare", "fakturanummer", "belopp" (Decimal), "source", "rad"}.
    Felrader rapporteras med källfil, radnummer och problem – de registreras
    inte som giltiga poster.
    """
    valid: list[dict] = []
    errors: list[dict] = []
    data_rows = 0

    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=";")
        try:
            header = next(reader)
        except StopIteration:
            raise ValueError(f"{path.name}: filen är tom (ingen rubrikrad).")

        header_map: dict[str, int] = {}
        for idx, name in enumerate(header):
            key = (name or "").strip().lower()
            if key and key not in header_map:
                header_map[key] = idx

        missing = [c for c in SOURCE_REQUIRED_COLUMNS if c.lower() not in header_map]
        if missing:
            raise ValueError(f"{path.name}: nödvändiga kolumner saknas: {', '.join(missing)}")

        col = {c: header_map[c.lower()] for c in SOURCE_REQUIRED_COLUMNS}

        def cell(row: list, name: str) -> str:
            i = col[name]
            return row[i] if i < len(row) else ""

        line_no = 1  # rubrikraden
        for row in reader:
            line_no += 1
            if not row or all(not (c or "").strip() for c in row):
                continue
            data_rows += 1

            raw_datum = cell(row, "Datum")
            raw_avsandare = cell(row, "Avsändare")
            raw_faktnr = cell(row, "Fakturanummer")
            raw_belopp = cell(row, "Belopp")

            problems: list[str] = []
            datum = parse_date(raw_datum)
            if datum is None:
                problems.append("ogiltigt eller saknat datum")
            fakturanummer = normalize_invoice(raw_faktnr)
            if not fakturanummer:
                problems.append("fakturanummer saknas")
            belopp = parse_amount(raw_belopp)
            if belopp is None:
                problems.append("ogiltigt eller saknat belopp")

            if problems:
                errors.append(
                    {
                        "kallfil": path.name,
                        "rad": line_no,
                        "problem": "; ".join(problems),
                        "datum": (raw_datum or "").strip(),
                        "avsandare": normalize_sender(raw_avsandare),
                        "fakturanummer": (raw_faktnr or "").strip(),
                        "belopp": (raw_belopp or "").strip(),
                    }
                )
                continue

            valid.append(
                {
                    "datum": datum,
                    "avsandare": normalize_sender(raw_avsandare),
                    "fakturanummer": fakturanummer,
                    "belopp": belopp,
                    "source": path.name,
                    "rad": line_no,
                }
            )

    return valid, errors, data_rows


# --- Inläsning av befintlig sammanställning -----------------------------------

def load_existing_workbook(path: Path) -> dict:
    """Läs befintlig sammanställning. Returnerar rows/errors/sources/meta.

    Avbryter med ValueError om filen inte följer förväntat format eller om en
    befintlig rad inte kan tolkas – hellre stopp än tyst dataförlust.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except PermissionError:
        raise ValueError(
            f"Kan inte läsa {path.name} – filen verkar vara öppen i Excel. "
            "Stäng filen och kör igen."
        )
    except Exception as exc:
        raise ValueError(f"Kunde inte öppna befintlig sammanställning: {exc}") from exc

    try:
        for sheet in (SHEET_FAKTUROR, SHEET_KALLFILER):
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"{path.name}: arket '{sheet}' saknas – filen följer inte "
                    "förväntat format och skrivs inte över. Flytta/byt namn på "
                    "filen om en ny sammanställning ska skapas."
                )

        rows: list[dict] = []
        ws = wb[SHEET_FAKTUROR]
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows_iter, [])]
        if header[: len(FAKTUROR_COLUMNS)] != FAKTUROR_COLUMNS:
            raise ValueError(
                f"{path.name}: kolumnerna i arket '{SHEET_FAKTUROR}' matchar inte "
                f"förväntade ({', '.join(FAKTUROR_COLUMNS)}). Filen skrivs inte över."
            )
        for i, row in enumerate(rows_iter, start=2):
            if row is None or all(v is None for v in row):
                continue
            datum = parse_date(row[0])
            fakturanummer = normalize_invoice(row[2])
            belopp = parse_amount(row[3])
            source = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            if datum is None or not fakturanummer or belopp is None or not source:
                raise ValueError(
                    f"{path.name}: rad {i} i arket '{SHEET_FAKTUROR}' kunde inte "
                    "tolkas (datum/fakturanummer/belopp/source). Kontrollera om "
                    "filen har redigerats manuellt – inget skrivs över."
                )
            rows.append(
                {
                    "datum": datum,
                    "avsandare": normalize_sender(row[1]),
                    "fakturanummer": fakturanummer,
                    "belopp": belopp,
                    "source": source,
                }
            )

        errors: list[dict] = []
        if SHEET_FEL in wb.sheetnames:
            ws = wb[SHEET_FEL]
            rows_iter = ws.iter_rows(values_only=True)
            next(rows_iter, None)  # rubrik
            for row in rows_iter:
                if row is None or all(v is None for v in row):
                    continue
                vals = list(row) + [None] * (len(FEL_COLUMNS) - len(row))
                errors.append(
                    {
                        "kallfil": str(vals[0] or "").strip(),
                        "rad": vals[1] if vals[1] is not None else "",
                        "problem": str(vals[2] or "").strip(),
                        "datum": str(vals[3] or "").strip(),
                        "avsandare": str(vals[4] or "").strip(),
                        "fakturanummer": normalize_invoice(vals[5]),
                        "belopp": str(vals[6] or "").strip(),
                    }
                )

        sources: dict[str, dict] = {}
        ws = wb[SHEET_KALLFILER]
        rows_iter = ws.iter_rows(values_only=True)
        next(rows_iter, None)  # rubrik
        for row in rows_iter:
            if row is None or all(v is None for v in row):
                continue
            vals = list(row) + [None] * (len(KALLFIL_COLUMNS) - len(row))
            name = str(vals[0] or "").strip()
            if not name:
                continue
            sources[name] = {
                "sha256": str(vals[1] or "").strip().lower(),
                "inlast": str(vals[2] or "").strip(),
                "period": str(vals[3] or "").strip(),
                "rader": vals[4],
                "giltiga": vals[5],
                "felrader": vals[6],
                "utanfor": vals[7],
            }

        meta: dict[str, str] = {}
        if SHEET_META in wb.sheetnames:
            ws = wb[SHEET_META]
            for row in ws.iter_rows(values_only=True):
                if row and row[0] is not None:
                    meta[str(row[0]).strip()] = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

        return {"rows": rows, "errors": errors, "sources": sources, "meta": meta}
    finally:
        wb.close()


# --- Skrivning av sammanställningen -------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
DUP_FILL = PatternFill("solid", fgColor="FFC7CE")
DUP_FONT = Font(color="9C0006")
ERR_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_workbook(
    target: Path,
    rows: list[dict],
    errors: list[dict],
    sources: dict[str, dict],
    meta: dict[str, str],
) -> None:
    """Skriv sammanställningen via temporär fil och atomiskt byte.

    Målfilen ersätts först när hela skrivningen lyckats; är den låst i Excel
    bevaras befintliga data och ett begripligt fel visas.
    """
    wb = openpyxl.Workbook()

    # --- Fakturor ---
    ws = wb.active
    ws.title = SHEET_FAKTUROR
    ws.append(FAKTUROR_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append(
            [r["datum"], r["avsandare"], r["fakturanummer"], float(r["belopp"]), r["status"], r["source"]]
        )
        i = ws.max_row
        ws.cell(row=i, column=3).number_format = "@"  # texten bevarar inledande nollor
        ws.cell(row=i, column=4).number_format = "#,##0.00"
        if r["status"] == STATUS_DUPLICATE:
            for c in range(1, len(FAKTUROR_COLUMNS) + 1):
                ws.cell(row=i, column=c).fill = DUP_FILL
                ws.cell(row=i, column=c).font = DUP_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FAKTUROR_COLUMNS))}{max(ws.max_row, 1)}"
    for col, width in zip("ABCDEF", (12, 34, 15, 13, 11, 46)):
        ws.column_dimensions[col].width = width

    # --- Felrader ---
    ws = wb.create_sheet(SHEET_FEL)
    ws.append(FEL_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for e in errors:
        ws.append(
            [e["kallfil"], e["rad"], e["problem"], e["datum"], e["avsandare"], e["fakturanummer"], e["belopp"]]
        )
        i = ws.max_row
        ws.cell(row=i, column=6).number_format = "@"
        for c in range(1, len(FEL_COLUMNS) + 1):
            ws.cell(row=i, column=c).fill = ERR_FILL
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFG", (46, 6, 44, 12, 30, 15, 13)):
        ws.column_dimensions[col].width = width

    # --- Källfiler (metadata för omkörningsskydd) ---
    ws = wb.create_sheet(SHEET_KALLFILER)
    ws.append(KALLFIL_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for name in sorted(sources):
        s = sources[name]
        ws.append(
            [name, s["sha256"], s["inlast"], s["period"], s["rader"], s["giltiga"], s["felrader"], s["utanfor"]]
        )
    for col, width in zip("ABCDEFGH", (46, 66, 20, 9, 7, 8, 9, 14)):
        ws.column_dimensions[col].width = width

    # --- Metadata ---
    ws = wb.create_sheet(SHEET_META)
    ws.append(["Nyckel", "Värde"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for key in sorted(meta):
        ws.append([key, meta[key]])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, target)
    except PermissionError:
        raise RuntimeError(
            f"Målfilen är låst (troligen öppen i Excel): {target}\n"
            "Stäng filen i Excel och kör skriptet igen. Befintliga data är oförändrade."
        )
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


# --- Hjälplogik ---------------------------------------------------------------

def resolve_csv_dir(raw: str) -> Path:
    """Tolka --input till mappen med registrerings-CSV:erna.

    Accepterar både output-mappen direkt och månadsmappen (då används dess
    undermapp "output").
    """
    cleaned = raw.strip().strip('"').strip("'").strip()
    path = Path(cleaned)
    if not path.is_dir():
        raise FileNotFoundError(f"Mappen finns inte: {path}")
    if find_source_files(path):
        return path
    sub = path / "output"
    if sub.is_dir() and find_source_files(sub):
        return sub
    # Ingen källfil någonstans – välj mest rimliga mappen för felmeddelandet.
    return sub if sub.is_dir() else path


def month_folder_name(csv_dir: Path) -> str:
    """Månadsmappens namn: output-mappens förälder (inte namnet 'output')."""
    if csv_dir.name.lower() == "output":
        return csv_dir.parent.name
    return csv_dir.name


def apply_duplicate_status(rows: list[dict]) -> int:
    """Sätt status på varje rad; returnerar antal dubblettgrupper.

    Samma fakturanummer i flera källposter => samtliga rader markeras
    "Dubblett" (även tidigare registrerade). Inga rader tas bort/slås ihop.
    """
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["fakturanummer"]] = counts.get(r["fakturanummer"], 0) + 1
    groups = 0
    for nr, n in counts.items():
        if n > 1:
            groups += 1
    for r in rows:
        r["status"] = STATUS_DUPLICATE if counts[r["fakturanummer"]] > 1 else STATUS_PREPARED
    return groups


def determine_month(arg_month: str | None, existing_period: str, new_rows: list[dict]) -> str:
    """Bestäm perioden (ÅÅÅÅ-MM).

    Prioritet: --month > befintlig fils period > entydig månad i nya rader.
    Perioden avser kalendermånaden i postens betalningsdatum – aldrig dagens
    datum.
    """
    if arg_month:
        if not MONTH_RE.match(arg_month):
            raise ValueError(f"Ogiltig --month: {arg_month!r} (förväntat format ÅÅÅÅ-MM).")
        if existing_period and existing_period != arg_month:
            raise ValueError(
                f"Befintlig sammanställning avser perioden {existing_period}, "
                f"men --month {arg_month} angavs. Varje fil innehåller endast en "
                "månad – kontrollera månadsmappen eller perioden."
            )
        return arg_month
    if existing_period:
        return existing_period
    months = sorted({r["datum"][:7] for r in new_rows})
    if not months:
        raise ValueError(
            "Ingen period kunde bestämmas: inga giltiga nya poster och ingen "
            "befintlig sammanställning. Ange --month ÅÅÅÅ-MM."
        )
    if len(months) > 1:
        per_month = {m: sum(1 for r in new_rows if r["datum"][:7] == m) for m in months}
        detail = ", ".join(f"{m} ({n} rader)" for m, n in per_month.items())
        raise ValueError(
            f"Källposterna spänner över flera månader: {detail}. "
            "Ange vilken månad som avses med --month ÅÅÅÅ-MM."
        )
    return months[0]


# --- Huvudprogram ---------------------------------------------------------------

def main() -> int:
    print("=" * 60)
    print(" Sammanställning: registrerings-CSV -> fakturor-lista (.xlsx)")
    print("=" * 60)

    parser = argparse.ArgumentParser(
        description="Sammanställ förberedda betalningsposter (Bankgiro + Swish) "
        "till en månadsvis kontrollista med dubblettmarkering."
    )
    parser.add_argument(
        "--input",
        type=str,
        default=None,
        help="Output-mappen med registrerings-CSV:erna, eller månadsmappen "
        "(då används dess undermapp 'output').",
    )
    parser.add_argument(
        "--month",
        type=str,
        default=None,
        help="Period ÅÅÅÅ-MM (kalendermånaden i betalningsdatum). Standard: "
        "befintlig fils period, annars entydig månad i källposterna.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Visa planerade ändringar utan att skriva några filer.",
    )
    parser.add_argument("path", nargs="?", default=None, help="Alternativ till --input.")
    args = parser.parse_args()

    raw_input_ = args.input if args.input else args.path
    if not raw_input_:
        print("FEL: Ange indata med --input <output-mapp eller månadsmapp>.", file=sys.stderr)
        return 1

    try:
        csv_dir = resolve_csv_dir(raw_input_)
    except FileNotFoundError as exc:
        print(f"FEL: {exc}", file=sys.stderr)
        return 1

    mappnamn = month_folder_name(csv_dir)
    target = csv_dir / TARGET_SUBDIR / f"{TARGET_PREFIX}_{mappnamn}.xlsx"

    print(f"Källmapp:    {csv_dir}")
    print(f"Månadsmapp:  {mappnamn}")
    print(f"Målfil:      {target}")
    if args.dry_run:
        print("Läge:        TORRKÖRNING – inga filer skrivs.")

    source_files = find_source_files(csv_dir)
    if not source_files:
        print(
            f"FEL: Inga källfiler ({' / '.join(p + '_*.csv' for p in SOURCE_PREFIXES)}) "
            f"hittades i {csv_dir}.",
            file=sys.stderr,
        )
        return 1

    # Läs befintlig sammanställning om den finns.
    existing = {"rows": [], "errors": [], "sources": {}, "meta": {}}
    if target.exists():
        try:
            existing = load_existing_workbook(target)
        except ValueError as exc:
            print(f"FEL: {exc}", file=sys.stderr)
            return 1
        print(f"Befintlig sammanställning inläst: {len(existing['rows'])} poster, "
              f"{len(existing['errors'])} felrader, {len(existing['sources'])} källfiler.")

    rows: list[dict] = list(existing["rows"])
    errors: list[dict] = list(existing["errors"])
    sources: dict[str, dict] = dict(existing["sources"])
    existing_period = existing["meta"].get("Period", "")

    # Gå igenom källfilerna: oförändrade hoppas över, ändrade rapporteras.
    new_rows: list[dict] = []
    new_errors: list[dict] = []
    skipped_unchanged: list[str] = []
    changed_files: list[str] = []
    pending: list[tuple[Path, str, list[dict], list[dict], int]] = []

    for path in source_files:
        digest = sha256_of_file(path)
        known = sources.get(path.name)
        if known is not None:
            if known["sha256"] == digest:
                skipped_unchanged.append(path.name)
            else:
                changed_files.append(path.name)
            continue
        try:
            valid, errs, total = read_source_csv(path)
        except (ValueError, OSError) as exc:
            print(f"FEL: {exc}", file=sys.stderr)
            return 1
        pending.append((path, digest, valid, errs, total))
        new_rows.extend(valid)
        new_errors.extend(errs)

    # Bestäm perioden och filtrera nya rader till kalendermånaden.
    try:
        period = determine_month(args.month, existing_period, new_rows)
    except ValueError as exc:
        print(f"FEL: {exc}", file=sys.stderr)
        return 1
    print(f"Period:      {period} (kalendermånaden i betalningsdatum)")

    out_of_period: dict[str, int] = {}
    kept_rows: list[dict] = []
    for r in new_rows:
        if r["datum"][:7] == period:
            kept_rows.append(r)
        else:
            out_of_period[r["source"]] = out_of_period.get(r["source"], 0) + 1

    # Registrera metadata för de nya källfilerna.
    now_str = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for path, digest, valid, errs, total in pending:
        kept = sum(1 for r in kept_rows if r["source"] == path.name)
        sources[path.name] = {
            "sha256": digest,
            "inlast": now_str,
            "period": period,
            "rader": total,
            "giltiga": kept,
            "felrader": len(errs),
            "utanfor": out_of_period.get(path.name, 0),
        }

    rows.extend(kept_rows)
    errors.extend(new_errors)

    # Dubblettstatus över samtliga rader (gamla + nya, Bankgiro + Swish).
    dup_groups = apply_duplicate_status(rows)
    rows.sort(key=lambda r: (r["datum"], r["fakturanummer"], r["source"], r["avsandare"]))

    total_amount = sum((r["belopp"] for r in rows), Decimal("0.00"))
    total_str = format_sum_sv(total_amount)

    meta = {
        "Period": period,
        "Månadsmapp": mappnamn,
        "Senast uppdaterad": now_str,
        "Antal poster": str(len(rows)),
        "Totalsumma (inkl. ev. dubbletter)": total_str,
        "Anmärkning": "Status 'Förberedd' = posten finns i importunderlaget; "
        "importen kan vara ogjord. Dubbletter ingår i totalsumman och "
        "räknas aldrig bort automatiskt.",
    }

    # --- Körningsrapport ---
    print("\nKörningsrapport")
    print(f"  Källfiler funna:          {len(source_files)}")
    print(f"  Nya källfiler inlästa:    {len(pending)}")
    print(f"  Redan inlästa (oförändrade, hoppas över): {len(skipped_unchanged)}")
    for name in skipped_unchanged:
        print(f"    - {name}")
    print(f"  Nya poster ({period}):    {len(kept_rows)}")
    if out_of_period:
        print(f"  Utanför perioden (ej medtagna): {sum(out_of_period.values())}")
        for name, n in sorted(out_of_period.items()):
            print(f"    - {name}: {n} rader")
    print(f"  Poster totalt i listan:   {len(rows)}")
    print(f"  Dubblettgrupper:          {dup_groups} "
          f"({sum(1 for r in rows if r['status'] == STATUS_DUPLICATE)} markerade rader)")
    print(f"  Felrader (nya):           {len(new_errors)}  (totalt {len(errors)})")
    for e in new_errors:
        print(f"    - {e['kallfil']} rad {e['rad']}: {e['problem']}")
    print(f"  Totalsumma:               {total_str}  (eventuella dubbletter ingår)")

    if changed_files:
        print("\n" + "!" * 60)
        print("VARNING: Följande källfiler har ÄNDRATS sedan de lästes in.")
        print("De har INTE lästs in igen och inga befintliga data har ändrats.")
        print("Kontrollera manuellt varför filen ändrats (SHA-256 skiljer sig):")
        for name in changed_files:
            print(f"  - {name}")
        print("!" * 60)

    if args.dry_run:
        print("\nTORRKÖRNING: inga filer skrevs.")
        if not target.exists():
            print(f"Skulle ha skapat: {target}")
        elif pending:
            print(f"Skulle ha uppdaterat: {target} (+{len(kept_rows)} poster, +{len(new_errors)} felrader)")
        else:
            print(f"Ingen förändring att skriva till: {target}")
        return 2 if changed_files else 0

    try:
        write_workbook(target, rows, errors, sources, meta)
    except RuntimeError as exc:
        print(f"\nFEL: {exc}", file=sys.stderr)
        return 1

    print(f"\nKlart! Sammanställningen sparad: {target}")
    return 2 if changed_files else 0


if __name__ == "__main__":
    raise SystemExit(main())
